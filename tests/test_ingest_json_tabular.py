"""Tests for deterministic JSON / CSV / XLSX ingestion (interpreter.ingest).

These are parser-grade, sport-agnostic extractors — the engine, not AI. They
must (a) turn each tabular format into TableCandidates feeding the existing
IngestStream contract, (b) keep XLSX disambiguated from generic ZIP (an .xlsx
IS a zip), (c) flow through _extract_zip so a mirror containing them "just
works" end-to-end via interpret_document, and (d) leave existing HTML/PDF/HY3/
ZIP sniffing unchanged.
"""

from __future__ import annotations

import io
import json
import zipfile

import pytest

from mediahub.interpreter import interpret_document
from mediahub.interpreter.ingest import _sniff_format, ingest


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------


def test_json_nested_results_array_becomes_table():
    blob = json.dumps(
        {
            "meet": "Spring Open 2026",
            "data": {
                "results": [
                    {"place": 1, "name": "Ada", "club": "Bton", "mark": "1:02.34"},
                    {"place": 2, "name": "Bea", "club": "Wgan", "mark": "1:03.11"},
                    {"place": 3, "name": "Cy", "club": "Hova", "mark": "1:04.50"},
                ]
            },
        }
    ).encode()
    stream = ingest(blob, content_type_hint="json")
    assert stream.format_detected == "json"
    assert len(stream.tables) == 1
    header = stream.tables[0].rows[0]
    assert "name" in header and "mark" in header
    assert ["1", "Ada", "Bton", "1:02.34"] == stream.tables[0].rows[1]
    assert "Spring Open 2026" in stream.text  # scalar metadata kept in text


def test_json_without_object_array_is_safe():
    stream = ingest(b'{"note": "no results here", "count": 0}', content_type_hint="json")
    assert stream.format_detected == "json"
    assert stream.tables == []  # nothing tabular, but no crash


def test_json_sniffed_without_hint():
    assert _sniff_format(b'  [{"a":1,"b":2,"c":3}]') == "json"
    assert _sniff_format(b'{"x":1}') == "json"


# ---------------------------------------------------------------------------
# CSV / TSV
# ---------------------------------------------------------------------------


def test_csv_ragged_rows_parse():
    csv_bytes = b"place,name,club,time\n1,Ada,Bton,1:02.3\n2,Bea\n3,Cy,Hova,1:04.5,extra\n"
    stream = ingest(csv_bytes, content_type_hint="csv")
    assert stream.format_detected == "csv"
    assert len(stream.tables) == 1
    rows = stream.tables[0].rows
    assert rows[0] == ["place", "name", "club", "time"]
    assert rows[2] == ["2", "Bea"]  # ragged row preserved
    assert rows[3][:4] == ["3", "Cy", "Hova", "1:04.5"]


def test_tsv_delimiter_sniffed():
    tsv = b"place\tname\ttime\n1\tAda\t58.21\n2\tBea\t59.10\n"
    stream = ingest(tsv, content_type_hint="tsv")
    assert stream.format_detected == "csv"
    assert stream.tables[0].rows[1] == ["1", "Ada", "58.21"]


def test_csv_sniffed_without_hint():
    assert _sniff_format(b"a,b,c\n1,2,3\n4,5,6\n") == "csv"
    # space-aligned text is NOT misread as CSV
    assert _sniff_format(b"1  Ada  58.21\n2  Bea  59.10\n") != "csv"


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_multisheet_one_table_per_sheet():
    xb = _make_xlsx(
        {
            "Heat 1": [["Place", "Name", "Time"], [1, "Ada", "1:02.34"], [2, "Bea", "1:03.11"]],
            "Heat 2": [["Place", "Name", "Time"], [1, "Cy", "58.21"]],
        }
    )
    stream = ingest(xb, content_type_hint="xlsx")
    assert stream.format_detected == "xlsx"
    assert len(stream.tables) == 2
    assert stream.tables[0].rows[1] == ["1", "Ada", "1:02.34"]


def test_xlsx_detected_before_generic_zip():
    """An .xlsx is a ZIP; sniffing MUST resolve it to xlsx, never zip."""
    xb = _make_xlsx({"S": [["a", "b", "c"], [1, 2, 3]]})
    assert xb[:4] == b"PK\x03\x04"  # it really is a zip container
    assert _sniff_format(xb) == "xlsx"  # but sniffed as xlsx, unhinted


# ---------------------------------------------------------------------------
# XLSX inside a mirror ZIP — end-to-end through interpret_document
# ---------------------------------------------------------------------------


def test_xlsx_inside_zip_flows_through_interpret_document():
    xb = _make_xlsx(
        {
            "Results": [
                ["Place", "Name", "YoB", "Club", "Time"],
                [1, "Ada Lovelace", 2009, "Brighton", "1:02.34"],
                [2, "Bea Carr", 2010, "Wigan", "1:03.11"],
            ]
        }
    )
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as zf:
        zf.writestr("results.xlsx", xb)
    zip_bytes = zbuf.getvalue()

    # ingest() proves the xlsx member's data survives into TableCandidates
    stream = ingest(zip_bytes)
    assert stream.format_detected == "zip"
    assert any("Ada Lovelace" in c for row in stream.tables[0].rows for c in row)
    assert "1:02.34" in stream.text

    # interpret_document runs the whole pipeline over the mirror ZIP without error
    meet = interpret_document(zip_bytes)
    assert meet is not None
    assert "format:zip" in meet.sources_used


# ---------------------------------------------------------------------------
# Regression — existing sniffing unchanged
# ---------------------------------------------------------------------------


def test_existing_sniffing_unchanged():
    assert _sniff_format(b"%PDF-1.7\n...") == "pdf"
    assert _sniff_format(b"<!DOCTYPE html><html><table></table></html>") == "html"
    assert _sniff_format(b"<html><body><table><tr><td>x</td></tr></table></body></html>") == "html"
    # hy3: many lines starting capital+digit
    assert _sniff_format(b"A1record\nB2record\nC3record\nD4record\nE5record\n") == "hy3"
    # a plain (non-xlsx) ZIP stays "zip"
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as zf:
        zf.writestr("meet.hy3", b"A1foo\nB2bar\n")
    assert _sniff_format(zbuf.getvalue()) == "zip"


# ---------------------------------------------------------------------------
# Mirror-sidecar skip + colliding-column resolution (results-from-a-link)
# ---------------------------------------------------------------------------


def test_mirror_sidecar_provenance_not_ingested_as_meet_name():
    """A results-fetch mirror ZIP carries a ``_provenance.json`` sidecar whose
    ``entry_url`` line otherwise leads the combined text stream and gets
    mistaken for the meet title. The sidecar (and screenshots) must be skipped
    so the real meet name from the result page wins."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        # Provenance written FIRST, mirroring results_fetch/package.py.
        z.writestr(
            "_provenance.json",
            json.dumps({"entry_url": "https://results.example.org/swimming/results/2026/champs/"}),
        )
        z.writestr("_screenshots/entry.png", b"\x89PNG not-a-real-image")
        z.writestr(
            "RG11H205.HTM",
            "<html><head><title>City Spring Meet 2026</title></head><body>"
            "<table><tr><th>Place</th><th>Name</th><th>Club</th><th>Time</th></tr>"
            "<tr><td>1.</td><td>Amy Crowley</td><td>City of Cardiff</td><td>33.45</td></tr>"
            "</table></body></html>",
        )
    res = interpret_document(buf.getvalue(), hint="zip")
    assert res.meet_name == "City Spring Meet 2026"
    assert "entry_url" not in (res.meet_name or "")
    assert "http" not in (res.meet_name or "")


def test_colliding_columns_keep_highest_confidence_value():
    """When two columns claim the same semantic type — a real time plus a lap
    split, or a place plus a separate points column — the highest-confidence
    column wins instead of the rightmost one overwriting the result."""
    from mediahub.interpreter.rows import _extract_swim_from_cells
    from mediahub.interpreter.schema_dataclasses import ColumnSchema

    schemas = [
        ColumnSchema(
            name="place", col_type="place", confidence=0.55, col_index=0, header_text="Place"
        ),
        ColumnSchema(
            name="name", col_type="name", confidence=0.90, col_index=1, header_text="Name"
        ),
        ColumnSchema(
            name="time", col_type="time", confidence=0.55, col_index=2, header_text="Time"
        ),
        ColumnSchema(name="time", col_type="time", confidence=0.35, col_index=3, header_text="50"),
        ColumnSchema(
            name="place", col_type="place", confidence=0.35, col_index=4, header_text="Pts"
        ),
    ]
    cells = ["1.", "Leila Kokas", "1:23.60", "37.91", "287"]
    swim = _extract_swim_from_cells(cells, schemas)
    assert swim is not None
    assert swim.swimmer_name == "Leila Kokas"
    assert swim.time == "1:23.60"  # the real time, not the 37.91 split
    assert swim.place == 1  # the real place ("1."), not the 287 points
